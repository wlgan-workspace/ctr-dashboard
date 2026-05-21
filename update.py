#!/usr/bin/env python3
"""
亲子频道 CTR 看板数据更新脚本
用法：python update.py [Excel文件路径]
     不传路径则自动在桌面寻找 亲子频道数据.xlsx
"""

import sys
import re
import json
import datetime
import subprocess
from pathlib import Path
from collections import defaultdict

REPO_DIR = Path(__file__).parent
DATA_JSON = REPO_DIR / "data.json"
DESKTOP = Path.home() / "Desktop"

DEFAULT_XLSX_NAMES = [
    "亲子频道数据.xlsx",
    "亲子频道数据 (2).xlsx",
    "亲子频道数据(2).xlsx",
]


# ── Excel → DATA ──────────────────────────────────────────────────────────────

def date_to_serial(d):
    """datetime.date → Excel serial number"""
    return (d - datetime.date(1899, 12, 30)).days


def serial_to_date(n):
    return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(n))


def read_rows(ws, min_row=2):
    return [r for r in ws.iter_rows(min_row=min_row, values_only=True) if r[0] is not None]


def parse_d(raw):
    if isinstance(raw, datetime.datetime):
        return date_to_serial(raw.date()), raw.strftime("%Y-%m-%d")
    n = int(raw)
    return n, serial_to_date(n).strftime("%Y-%m-%d")


def load_component_meta(ws):
    """
    Sheet 4 (组件埋点): col0=中文名, col1=英文logkey
    返回 (componentMeta list, logkey→(type, comp_name, rule) dict)
    """
    PREFIX = "亲子频道_首页_"
    SUFFIX_EXPOSURE = ("_曝光", "_exposure")
    SUFFIX_CLICK = ("_点击", "_click")

    comps = {}   # comp_name → dict
    order_map = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        cn_raw, en_raw = row[0], row[1]
        if not cn_raw or not en_raw:
            continue
        cn = str(cn_raw).strip().rstrip("\n")
        en = str(en_raw).strip().rstrip("\n")

        name = cn[len(PREFIX):] if cn.startswith(PREFIX) else cn

        is_exposure = name.endswith("_曝光")
        is_click = name.endswith("_点击")
        # also handle 页面UV pseudo-entries
        if name == "页面UV":
            continue

        if is_exposure:
            comp_name = name[:-3]
        elif is_click:
            comp_name = name[:-3]
        else:
            comp_name = name

        if comp_name not in comps:
            comps[comp_name] = {}
            order_map[comp_name] = len(order_map) + 1

        if is_exposure:
            comps[comp_name]["exposureKey"] = en
            comps[comp_name]["exposureCn"] = cn
        elif is_click:
            comps[comp_name]["clickKey"] = en
            comps[comp_name]["clickCn"] = cn

    meta = []
    logkey_map = {}  # logkey → (kind, comp_name, rule)
    for comp_name, info in comps.items():
        exp_key = info.get("exposureKey", "")
        clk_key = info.get("clickKey", "")
        rule = "组件曝光UV" if exp_key else "页面UV"
        entry = {
            "component": comp_name,
            "order": order_map[comp_name],
            "exposureKey": exp_key,
            "clickKey": clk_key,
            "exposureRule": rule,
            "exposureCn": info.get("exposureCn", "页面UV"),
            "clickCn": info.get("clickCn", ""),
            "exposureKeyFound": bool(exp_key),
            "clickKeyFound": bool(clk_key),
        }
        meta.append(entry)
        if exp_key:
            logkey_map[exp_key] = ("exposure", comp_name, rule)
        if clk_key:
            logkey_map[clk_key] = ("click", comp_name, rule)

    meta.sort(key=lambda x: x["order"])
    return meta, logkey_map


def build_data(xlsx_path):
    try:
        import openpyxl
    except ImportError:
        print("缺少 openpyxl，正在安装…")
        subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl"], check=True)
        import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    sheets = wb.worksheets

    # ── Sheet 0: 页面数据 (d, region, locale, uv, avg_duration) ──
    page_map = {}
    for row in read_rows(sheets[0]):
        d_num, date_str = parse_d(row[0])
        region, locale, uv = str(row[1]), str(row[2]), int(row[3] or 0)
        key = (d_num, region, locale)
        page_map[key] = {
            "d": d_num, "date": date_str,
            "region": region, "locale": locale,
            "pageUV": uv, "flowUV": 0,
        }

    # ── Sheet 1: 导流页面数据 (d, region, locale, uv) ──
    for row in read_rows(sheets[1]):
        d_num, _ = parse_d(row[0])
        key = (d_num, str(row[1]), str(row[2]))
        if key in page_map:
            page_map[key]["flowUV"] = int(row[3] or 0)

    pageRows = sorted(page_map.values(), key=lambda x: (x["d"], x["region"], x["locale"]))

    # ── Sheet 3: 组件埋点 (component meta) ──
    comp_meta, logkey_map = load_component_meta(sheets[3])

    # ── Sheet 2: 组件数据 (d, region, locale, logkey, uv) ──
    comp_data = {}  # (d, region, locale, comp_name) → {exposureUV, clickUV, ...}
    for row in read_rows(sheets[2]):
        d_num, date_str = parse_d(row[0])
        region, locale, logkey = str(row[1]), str(row[2]), str(row[3]).strip().rstrip("\n")
        uv = int(row[4] or 0)
        if logkey not in logkey_map:
            continue
        kind, comp_name, rule = logkey_map[logkey]
        key = (d_num, region, locale, comp_name)
        if key not in comp_data:
            comp_data[key] = {
                "d": d_num, "date": date_str,
                "region": region, "locale": locale,
                "component": comp_name,
                "exposureUV": 0, "clickUV": 0,
                "exposureRule": rule,
            }
        if kind == "exposure":
            comp_data[key]["exposureUV"] = uv
        else:
            comp_data[key]["clickUV"] = uv

    # For components with 页面UV rule, fill exposureUV from page_map
    page_by_key = {(r["d"], r["region"], r["locale"]): r for r in pageRows}
    page_uv_comps = {m["component"] for m in comp_meta if m["exposureRule"] == "页面UV"}

    for key, row in list(comp_data.items()):
        if row["component"] in page_uv_comps:
            pk = (key[0], key[1], key[2])
            row["exposureUV"] = page_by_key.get(pk, {}).get("pageUV", 0)

    # For 页面UV components, also ensure rows exist for every (d, region, locale)
    for m in comp_meta:
        if m["exposureRule"] != "页面UV":
            continue
        for pr in pageRows:
            key = (pr["d"], pr["region"], pr["locale"], m["component"])
            if key not in comp_data:
                comp_data[key] = {
                    "d": pr["d"], "date": pr["date"],
                    "region": pr["region"], "locale": pr["locale"],
                    "component": m["component"],
                    "exposureUV": pr["pageUV"],
                    "clickUV": 0,
                    "exposureRule": "页面UV",
                }

    componentRows = sorted(comp_data.values(),
                           key=lambda x: (x["d"], x["region"], x["locale"], x["component"]))

    # ── Summary ──
    summary = _compute_summary(pageRows, componentRows, comp_meta)
    summary["updateTime"] = datetime.date.today().isoformat()

    missing = [m["exposureKey"] for m in comp_meta if not m["exposureKeyFound"] and m["exposureKey"]]
    missing += [m["clickKey"] for m in comp_meta if not m["clickKeyFound"] and m["clickKey"]]
    summary["missingLogkeys"] = missing

    return {
        "pageRows": pageRows,
        "componentRows": componentRows,
        "componentMeta": comp_meta,
        "summary": summary,
    }


def _avg(values):
    vals = [v for v in values if v is not None]
    return sum(vals) / len(vals) if vals else None


def _compute_summary(pageRows, componentRows, comp_meta):
    dates = sorted({r["date"] for r in pageRows})
    regions = sorted({r["region"] for r in pageRows})
    locales = sorted({r["locale"] for r in pageRows})

    # Overall page metrics
    total_page_uv = sum(r["pageUV"] for r in pageRows)
    total_flow_uv = sum(r["flowUV"] for r in pageRows)
    days = len(dates)

    # avgFlowRate = mean of daily (flowUV/pageUV) per locale row
    flow_rates = []
    for r in pageRows:
        if r["pageUV"] > 0:
            flow_rates.append(r["flowUV"] / r["pageUV"])

    # avgPageCTR — compute from component rows: mean of all component CTRs?
    # Or page-level: average flowUV/pageUV * 100? Looking at 4.3% ≈ some component avg
    # Use: mean of all non-null CTRs across componentRows (where exposureUV>0)
    all_ctrs = []
    for r in componentRows:
        if r["exposureUV"] > 0:
            all_ctrs.append(r["clickUV"] / r["exposureUV"] * 100)

    overall = {
        "avgPageCTR": _avg(all_ctrs),
        "avgFlowRate": _avg(flow_rates),
        "pageUV": total_page_uv,
        "flowUV": total_flow_uv,
        "days": days,
    }

    # Per-region page metrics
    region_page = []
    for reg in regions:
        r_rows = [r for r in pageRows if r["region"] == reg]
        r_comp = [r for r in componentRows if r["region"] == reg]
        ctrs = [r["clickUV"]/r["exposureUV"]*100 for r in r_comp if r["exposureUV"] > 0]
        fr = [r["flowUV"]/r["pageUV"] for r in r_rows if r["pageUV"] > 0]
        region_page.append({
            "region": reg,
            "avgPageCTR": _avg(ctrs),
            "avgFlowRate": _avg(fr),
            "pageUV": sum(r["pageUV"] for r in r_rows),
            "flowUV": sum(r["flowUV"] for r in r_rows),
            "days": len({r["date"] for r in r_rows}),
        })
    region_page.sort(key=lambda x: x["avgPageCTR"] or 0, reverse=True)

    # Component overall
    comp_names = [m["component"] for m in comp_meta]
    comp_overall = []
    for comp in comp_names:
        c_rows = [r for r in componentRows if r["component"] == comp]
        exp_uv = sum(r["exposureUV"] for r in c_rows)
        clk_uv = sum(r["clickUV"] for r in c_rows)
        ctrs = [r["clickUV"]/r["exposureUV"] for r in c_rows if r["exposureUV"] > 0]
        comp_overall.append({
            "component": comp,
            "avgCTR": _avg(ctrs),
            "exposureUV": exp_uv,
            "clickUV": clk_uv,
            "days": len({r["date"] for r in c_rows}),
        })
    comp_overall.sort(key=lambda x: x["avgCTR"] or 0, reverse=True)

    # Region × Component
    region_comp = []
    for reg in regions:
        for comp in comp_names:
            c_rows = [r for r in componentRows if r["region"] == reg and r["component"] == comp]
            if not c_rows:
                continue
            exp_uv = sum(r["exposureUV"] for r in c_rows)
            clk_uv = sum(r["clickUV"] for r in c_rows)
            ctrs = [r["clickUV"]/r["exposureUV"] for r in c_rows if r["exposureUV"] > 0]
            region_comp.append({
                "region": reg,
                "component": comp,
                "avgCTR": _avg(ctrs),
                "exposureUV": exp_uv,
                "clickUV": clk_uv,
                "days": len({r["date"] for r in c_rows}),
            })

    return {
        "dateMin": dates[0] if dates else "",
        "dateMax": dates[-1] if dates else "",
        "regions": regions,
        "locales": locales,
        "componentCount": len(comp_meta),
        "recordCounts": {"pageRows": len(pageRows), "componentRows": len(componentRows)},
        "overallPage": overall,
        "regionPage": region_page,
        "componentOverall": comp_overall,
        "regionComponent": region_comp,
    }


# ── Git push ───────────────────────────────────────────────────────────────────

def git_push(data_json_path, xlsx_name):
    repo = data_json_path.parent
    today = datetime.date.today().isoformat()
    msg = f"Update data: {xlsx_name} ({today})"

    def run(cmd):
        result = subprocess.run(cmd, cwd=repo, capture_output=True, text=True)
        if result.returncode != 0:
            print(f"  ✗ {' '.join(cmd)}")
            print(result.stderr.strip())
            return False
        return True

    print("\n📤 推送到 GitHub…")
    ok = (
        run(["git", "add", "data.json"])
        and run(["git", "commit", "-m", msg])
        and run(["git", "push"])
    )
    if ok:
        print(f"  ✓ 已推送：{msg}")
    else:
        print("  ✗ 推送失败，请检查 git 配置或网络")
    return ok


# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    # Determine xlsx path
    if len(sys.argv) > 1:
        xlsx_path = Path(sys.argv[1])
    else:
        xlsx_path = None
        for name in DEFAULT_XLSX_NAMES:
            p = DESKTOP / name
            if p.exists():
                xlsx_path = p
                break
        if xlsx_path is None:
            print("❌ 未找到 Excel 文件，请将文件放到桌面并命名为：亲子频道数据.xlsx")
            print("   或直接传入路径：python update.py 路径/文件.xlsx")
            sys.exit(1)

    print(f"📂 读取数据：{xlsx_path.name}")

    # Build DATA
    data = build_data(xlsx_path)
    print(f"  pageRows: {len(data['pageRows'])}  componentRows: {len(data['componentRows'])}")
    print(f"  日期范围: {data['summary']['dateMin']} ~ {data['summary']['dateMax']}")

    # Write data.json
    with open(DATA_JSON, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print(f"✅ data.json 已更新（{DATA_JSON.stat().st_size // 1024} KB）")

    # Git push
    git_push(DATA_JSON, xlsx_path.name)


if __name__ == "__main__":
    main()
